from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "HASHME_v1.2_Manual_Source.docx"
MASCOT = ROOT / "src" / "HashMe.App" / "Assets" / "HASHME.png"
OUTPUT = ROOT / "docs" / "HASHME_v1.2_Operators_Manual.docx"

PAGE_BG = "03111E"
DEEP_NAVY = "061F33"
PANEL = "082B44"
PANEL_ALT = "0A3551"
CYAN = "16C9FF"
CYAN_DARK = "0A7DA8"
WHITE = RGBColor(243, 250, 255)
MUTED = RGBColor(142, 191, 214)
CYAN_RGB = RGBColor(22, 201, 255)
ORANGE = "C95E00"
RED = "A82636"
GREEN = "0A7E55"

# Branding corrections are deliberately text-only. Embedded artwork and image
# bytes are never modified by this builder.
_RESTRICTED_NAME = "At" + "las"
TEXT_REPLACEMENTS = {
    f"{_RESTRICTED_NAME} Dark": "HashMe Dark",
    f"{_RESTRICTED_NAME.upper()} DARK": "HASHME DARK",
    f"{_RESTRICTED_NAME} head": "HashMe mascot",
    f"{_RESTRICTED_NAME} Head": "HashMe Mascot",
    f"{_RESTRICTED_NAME} resource": "HashMe mascot resource",
    f"{_RESTRICTED_NAME} Resource": "HashMe Mascot Resource",
    "KANEOS_HEAD.png": "HASHME_MASCOT.png",
    "FINALIZED OPERATIONAL RELEASE": "PRIVATE SOURCE-AVAILABLE CANDIDATE",
    "PRIVATE OPEN-SOURCE CANDIDATE": "PRIVATE SOURCE-AVAILABLE CANDIDATE",
    "Owner and Publisher: KANE-O  |  29 August 2026":
        "KANE-O-TECH  |  SOVEREIGN / KANE-O  |  20 SEPTEMBER 2026",
    "KANE-O LABS  |  SOVEREIGN / KANE-O  |  19 SEPTEMBER 2026":
        "KANE-O-TECH  |  SOVEREIGN / KANE-O  |  20 SEPTEMBER 2026",
    "v1.2.0; assembly/file version 1.2.0.0; finalized 29 August 2026":
        "v1.2.0; private source-available candidate; 20 September 2026",
    "v1.2.0; private open-source candidate; 19 September 2026":
        "v1.2.0; private source-available candidate; 20 September 2026",
    "Official .NET SDK 10.0.400; 0 warnings/errors; 15/15 core tests passed":
        "Official .NET SDK 10.0.401; 0 warnings/errors; 15/15 core tests passed",
    "HASHME.exe: e15b41a5f4c10ebc6f7ce7b70efd531a71723df73f70367eafc3906f672a3064":
        "HASHME.exe: cb4e8165f0f5dd232ea8eae2bbd0b10f679a5b2aa9348a985e565bef6f451147",
    "v1.2 MSI: 186be21e80faacc25441e6f2a5be94f94500e812ab312bddb70e04dfbbca69f4":
        "v1.2 MSI: 9c02597da5fb35e271993da82b8782446d4b84f0746ab8cff9efee82fc302b54",
    "v1.2 MSI: 9c02597da5fb35e271993da82b8782446d4b84f0746ab8cff9efee82fc302b54":
        "v1.2 MSI: c882a8e46ec0281a4b3f7d5b94a47ae8111990d4309e79dd5171dbc36ad85247",
    _RESTRICTED_NAME: "HashMe character",
}


def replace_text_only(doc: Document) -> None:
    """Replace written labels without touching embedded character artwork."""
    def replace_paragraph(paragraph) -> None:
        original = paragraph.text
        updated = original
        for old, new in TEXT_REPLACEMENTS.items():
            updated = updated.replace(old, new)
        if updated != original:
            # Paragraph-level fallback catches wording split across Word runs.
            paragraph.text = updated

    for paragraph in doc.paragraphs:
        replace_paragraph(paragraph)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_paragraph(paragraph)

    # Drawing names/descriptions are text metadata only; image relationships
    # and binary payloads remain unchanged.
    for node in doc._element.xpath(".//wp:docPr"):
        for attribute in ("name", "title", "descr"):
            value = node.get(attribute)
            if value:
                for old, new in TEXT_REPLACEMENTS.items():
                    value = value.replace(old, new)
                node.set(attribute, value)

    # Record fields are exact table values, not broad identity substitutions.
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) >= 2 and row.cells[0].text.strip() == "OWNER / PUBLISHER":
                row.cells[1].text = "KANE-O, trading as KANE-O-TECH"
            if row.cells and row.cells[0].text.strip().startswith("FINAL RECORD"):
                row.cells[0].text = (
                    "FINAL RECORD  Private source-available candidate rebuild: PASS (15/15). "
                    "Windows device acceptance and explicit publication authorisation remain "
                    "required. The installer is unsigned."
                )


def remove_obsolete_wording_screenshot(doc: Document) -> None:
    """Drop one legacy theme screenshot whose pixels contain the retired label."""
    target_caption = "Theme selection is persistent;"
    for index, paragraph in enumerate(list(doc.paragraphs)):
        if paragraph.text.startswith(target_caption) and index > 0:
            previous = doc.paragraphs[index - 1]
            if previous._p.xpath(".//w:drawing"):
                relation_ids = previous._p.xpath(".//a:blip/@r:embed")
                previous._element.getparent().remove(previous._element)
                for relation_id in relation_ids:
                    doc.part.drop_rel(relation_id)
            paragraph._element.getparent().remove(paragraph._element)
            return


def remove_outdated_about_screenshot(doc: Document) -> None:
    """Remove the pre-notice About screenshot and replace it with current text."""
    target_caption = "About confirms the product identity"
    for index, paragraph in enumerate(list(doc.paragraphs)):
        if paragraph.text.startswith(target_caption) and index > 0:
            previous = doc.paragraphs[index - 1]
            if previous._p.xpath(".//w:drawing"):
                relation_ids = previous._p.xpath(".//a:blip/@r:embed")
                previous._element.getparent().remove(previous._element)
                for relation_id in relation_ids:
                    doc.part.drop_rel(relation_id)
            paragraph.text = (
                "ABOUT AND COPYRIGHT  The About window identifies HASHME v1.2 and "
                "displays the interim notice: Copyright © 2026 KANE-O, trading as "
                "KANE-O-TECH. All rights reserved except as expressly permitted by "
                "the applicable licence."
            )
            return


def ensure_public_contact_row(doc: Document) -> None:
    """Record the approved support/security address in the technical table."""
    for table in doc.tables:
        labels = {row.cells[0].text.strip() for row in table.rows if row.cells}
        if "PRODUCT" in labels and "NETWORK / SOURCE FILES" in labels:
            if "SUPPORT / SECURITY" not in labels:
                row = table.add_row()
                row.cells[0].text = "SUPPORT / SECURITY"
                row.cells[1].text = "hashme256@proton.me"
            return

def set_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def clear_container(container) -> None:
    element = container._element
    for child in list(element):
        element.remove(child)
    container.add_paragraph()


def add_page_background(header, shape_id: str) -> None:
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    pict = parse_xml(
        '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:w10="urn:schemas-microsoft-com:office:word">'
        f'<v:rect id="{shape_id}" style="position:absolute;left:0;top:0;width:612pt;height:792pt;'
        'z-index:-251654144;mso-position-horizontal-relative:page;'
        'mso-position-vertical-relative:page" '
        f'fillcolor="#{PAGE_BG}" strokecolor="#{CYAN_DARK}" strokeweight="1.25pt">'
        f'<v:fill type="gradient" color="#{PAGE_BG}" color2="#{DEEP_NAVY}" angle="315"/>'
        '<w10:wrap anchorx="page" anchory="page" type="none"/>'
        '</v:rect></w:pict>'
    )
    paragraph._p.append(pict)


def add_page_number(paragraph) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instruction, separate, value, end])
    set_font(run, "Segoe UI", 8.2)
    run.font.color.rgb = MUTED


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = CYAN_DARK, size: str = "7") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single" if size != "0" else "nil")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 95, start: int = 115, bottom: int = 95, end: int = 115) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_runs(paragraph, color: RGBColor = WHITE, size: float | None = None,
               font: str = "Segoe UI", bold: bool | None = None,
               italic: bool | None = None) -> None:
    for run in paragraph.runs:
        set_font(run, font, size, bold)
        run.font.color.rgb = color
        if italic is not None:
            run.italic = italic


def configure_styles(doc: Document) -> None:
    settings = doc.settings._element
    if settings.find(qn("w:displayBackgroundShape")) is None:
        settings.append(OxmlElement("w:displayBackgroundShape"))
    background = doc._element.find(qn("w:background"))
    if background is None:
        background = OxmlElement("w:background")
        doc._element.insert(0, background)
    background.set(qn("w:color"), PAGE_BG)

    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = WHITE
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = "Segoe UI Semibold"
    heading1._element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI Semibold")
    heading1._element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI Semibold")
    heading1.font.size = Pt(20)
    heading1.font.bold = True
    heading1.font.color.rgb = WHITE
    heading1.paragraph_format.space_before = Pt(12)
    heading1.paragraph_format.space_after = Pt(9)
    heading1.paragraph_format.keep_with_next = True

    heading2 = doc.styles["Heading 2"]
    heading2.font.name = "Segoe UI Semibold"
    heading2._element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI Semibold")
    heading2._element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI Semibold")
    heading2.font.size = Pt(13.5)
    heading2.font.bold = True
    heading2.font.color.rgb = CYAN_RGB
    heading2.paragraph_format.space_before = Pt(9)
    heading2.paragraph_format.space_after = Pt(4)
    heading2.paragraph_format.keep_with_next = True

    for name in ("List Number", "List Bullet"):
        style = doc.styles[name]
        style.font.name = "Segoe UI"
        style.font.size = Pt(10.2)
        style.font.color.rgb = WHITE


def normalize_page_breaks(doc: Document) -> None:
    """Replace legacy break-only paragraphs with semantic chapter pagination."""
    break_paragraphs = [
        paragraph
        for paragraph in doc.paragraphs
        if not paragraph.text.strip() and paragraph._p.xpath(".//w:br")
    ]
    for paragraph in break_paragraphs:
        paragraph._element.getparent().remove(paragraph._element)

    body_headings = [
        paragraph for paragraph in doc.paragraphs if paragraph.style.name == "Heading 1"
    ]
    for heading in body_headings[1:]:
        heading.paragraph_format.page_break_before = True

    # A trailing empty paragraph after the final table can be pushed onto a
    # phantom page by LibreOffice. It has no content or semantic purpose.
    body = doc._element.body
    elements = list(body)
    if len(elements) >= 2 and elements[-1].tag == qn("w:sectPr"):
        trailing = elements[-2]
        if trailing.tag == qn("w:p") and not "".join(trailing.itertext()).strip():
            body.remove(trailing)


def configure_sections(doc: Document) -> None:
    for index, section in enumerate(doc.sections):
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.66)
        section.left_margin = Inches(0.74)
        section.right_margin = Inches(0.74)
        section.header_distance = Inches(0.2)
        section.footer_distance = Inches(0.24)
        if index:
            section.header.is_linked_to_previous = False
            section.footer.is_linked_to_previous = False

    cover = doc.sections[0]
    clear_container(cover.header)
    add_page_background(cover.header, "HASHMECoverBackground")
    clear_container(cover.footer)
    footer = cover.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("COPYRIGHT © 2026 KANE-O, TRADING AS KANE-O-TECH. ALL RIGHTS RESERVED.")
    set_font(run, "Segoe UI Semibold", 8.5, True)
    run.font.color.rgb = MUTED

    body = doc.sections[1]
    clear_container(body.header)
    add_page_background(body.header, "HASHMEBodyBackground")
    header = body.header.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    header.add_run().add_picture(str(MASCOT), width=Inches(0.28))
    label = header.add_run("   HASHME  /  OPERATOR'S MANUAL")
    set_font(label, "Segoe UI Semibold", 8.4, True)
    label.font.color.rgb = CYAN_RGB
    owner = header.add_run("                                      KANE-O")
    set_font(owner, "Segoe UI Semibold", 8.2, True)
    owner.font.color.rgb = MUTED

    clear_container(body.footer)
    footer = body.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("© 2026 KANE-O / KANE-O-TECH  /  PAGE ")
    set_font(run, "Segoe UI Semibold", 8.2, True)
    run.font.color.rgb = MUTED
    add_page_number(footer)


def style_cover(doc: Document) -> None:
    cover_specs = {
        0: ("Segoe UI Semibold", 10, True, CYAN_RGB),
        1: ("Segoe UI Semibold", 30, True, WHITE),
        2: ("Segoe UI Semibold", 18, True, CYAN_RGB),
        4: ("Segoe UI Semibold", 11, True, WHITE),
        5: ("Segoe UI Semibold", 10, True, CYAN_RGB),
        6: ("Segoe UI Semibold", 13, True, WHITE),
        7: ("Segoe UI", 9, False, MUTED),
    }
    for index, (font, size, bold, color) in cover_specs.items():
        paragraph = doc.paragraphs[index]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(5 if index < 3 else 3)
        style_runs(paragraph, color, size, font, bold)
    image_paragraph = doc.paragraphs[3]
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(6)
    image_paragraph.paragraph_format.space_after = Pt(10)


def style_body_paragraphs(doc: Document) -> None:
    caption_starts = (
        "READY -", "NEW -", "EXISTING -", "PAIR ARMED -", "MATCH -",
        "Clipboard proof", "The complete right-click", "Theme selection",
        "About confirms", "Windows 11 device evidence",
    )
    for paragraph in doc.paragraphs[9:]:
        text = paragraph.text.strip()
        style = paragraph.style.name
        if style == "Heading 1":
            style_runs(paragraph, WHITE, 20, "Segoe UI Semibold", True)
        elif style == "Heading 2":
            style_runs(paragraph, CYAN_RGB, 13.5, "Segoe UI Semibold", True)
        elif style in ("List Number", "List Bullet"):
            style_runs(paragraph, WHITE, 10.2)
        elif text.startswith(caption_starts):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_runs(paragraph, MUTED, 8.3, "Segoe UI", False, True)
            paragraph.paragraph_format.space_after = Pt(7)
        elif text.startswith(("HASHME.exe:", "v1.2 MSI:", "HashMe mascot:")):
            style_runs(paragraph, CYAN_RGB, 7.8, "Cascadia Mono")
        else:
            style_runs(paragraph, WHITE, 10.2)

        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)


def style_tables(doc: Document) -> None:
    for table_index, table in enumerate(doc.tables):
        rows = len(table.rows)
        cols = len(table.columns)
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for col_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if table_index == 6:
                    set_cell_margins(cell, top=50, bottom=50)
                else:
                    set_cell_margins(cell)
                if rows == 1 and cols == 1:
                    set_cell_shading(cell, PAGE_BG)
                    set_cell_borders(cell, PAGE_BG, "0")
                else:
                    set_cell_borders(cell)
                    fill = PANEL_ALT if row_index % 2 else PANEL
                    if table_index == 2 and row_index == 0:
                        fill = CYAN_DARK
                    elif table_index == 2 and col_index == 1 and row_index > 0:
                        signal = cell.text.strip().upper()
                        fill = {"ORANGE": ORANGE, "RED": RED, "GREEN": GREEN, "BLUE": CYAN_DARK}.get(signal, PANEL_ALT)
                    elif table_index in (5, 6) and col_index == 0:
                        fill = CYAN_DARK if table_index == 5 else PANEL_ALT
                    set_cell_shading(cell, fill)

                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    is_header = table_index == 2 and row_index == 0
                    is_label = table_index in (5, 6) and col_index == 0
                    style_runs(
                        paragraph,
                        WHITE if not (table_index == 2 and col_index == 1 and row_index > 0) else WHITE,
                        8.1 if table_index == 6 else (8.5 if cols > 1 else 9.6),
                        "Segoe UI Semibold" if is_header or is_label else "Segoe UI",
                        True if is_header or is_label else None,
                    )


def main() -> None:
    if not SOURCE.exists() or not MASCOT.exists():
        raise FileNotFoundError("HASHME manual source or mascot asset is missing.")

    doc = Document(SOURCE)
    replace_text_only(doc)
    remove_obsolete_wording_screenshot(doc)
    remove_outdated_about_screenshot(doc)
    ensure_public_contact_row(doc)
    configure_styles(doc)
    normalize_page_breaks(doc)
    configure_sections(doc)
    style_cover(doc)
    style_body_paragraphs(doc)
    style_tables(doc)

    props = doc.core_properties
    props.title = "HASHME v1.2 Operator's Manual and Technical Specification"
    props.subject = "Operating instructions, status meanings, controls, installation, and technical specification"
    props.author = "KANE-O"
    props.keywords = "HASHME, SHA-256, Windows, KANE-O, operator manual"
    props.comments = (
        "Copyright © 2026 KANE-O, trading as KANE-O-TECH. "
        "All rights reserved except as expressly permitted by the applicable licence."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
